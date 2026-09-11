"""Cau tra loi di ra Excel va Word, va khong duoc rung gi doc duong.

Ba thu phai song sot: canh bao do tin cay, metric key, va nhung gi khong xac
lap duoc. Mot bang Excel toan so ma bo mat dong "nhom nay chi co 3 dong" thi
nguy hiem hon la khong co bang.
"""

from __future__ import annotations

from io import BytesIO

from analysis_system.contracts.agents import ClaimEvidence, DataNeed, ManagerAnswer
from analysis_system.services.export_answer import SHEET_CLAIMS, to_excel, to_word

TRA_LOI = ManagerAnswer(
    question="Kênh thông tin nào được dùng nhiều nhất?",
    claims=(
        ClaimEvidence(
            claim="Tư vấn tài chính là kênh phổ biến nhất với 40 % số người trả lời.",
            metric_keys=("Source.Financial_Consultants.share_pct",),
            evidence_ref="mart://finance.parquet",
        ),
        ClaimEvidence(
            claim="Internet là kênh ít được dùng nhất với 12.50 %.",
            metric_keys=("Source.Internet.share_pct",),
            evidence_ref="mart://finance.parquet",
        ),
    ),
    warnings=("Debentures theo Avenue: bỏ qua 1 nhóm có dưới 5 dòng, quá ít để nói gì",),
    unanswered=("Không tự chạy hồi quy, chọn biến giải thích là một nhận định.",),
    needs=(
        DataNeed(ask="Cần cột ngày để nói được xu hướng.", blocked_by="khong co cot thoi gian"),
    ),
)

TRONG = ManagerAnswer(question="Câu hỏi không trả lời được")


def _excel_text(answer: ManagerAnswer) -> str:
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(to_excel(answer)))
    found = []
    for sheet in book.worksheets:
        found.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            found.extend(str(cell) for cell in row if cell is not None)
    return "\n".join(found)


def _word_text(answer: ManagerAnswer) -> str:
    from docx import Document

    document = Document(BytesIO(to_word(answer)))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


# --- khong duoc rung gi doc duong ---------------------------------------------


def test_excel_carries_the_reliability_warnings() -> None:
    assert "quá ít để nói gì" in _excel_text(TRA_LOI)


def test_word_carries_the_reliability_warnings() -> None:
    assert "quá ít để nói gì" in _word_text(TRA_LOI)


def test_excel_carries_the_metric_keys() -> None:
    """Khong lan nguoc duoc ve chi so thi ket luan chi la mot y kien."""
    assert "Source.Financial_Consultants.share_pct" in _excel_text(TRA_LOI)


def test_word_carries_the_metric_keys() -> None:
    assert "Source.Financial_Consultants.share_pct" in _word_text(TRA_LOI)


def test_excel_carries_what_could_not_be_established() -> None:
    assert "chọn biến giải thích là một nhận định" in _excel_text(TRA_LOI)


def test_word_carries_what_could_not_be_established() -> None:
    assert "chọn biến giải thích là một nhận định" in _word_text(TRA_LOI)


def test_every_claim_travels() -> None:
    text = _excel_text(TRA_LOI)
    for claim in TRA_LOI.claims:
        assert claim.claim in text


def test_the_question_travels() -> None:
    assert TRA_LOI.question in _excel_text(TRA_LOI)
    assert TRA_LOI.question in _word_text(TRA_LOI)


def test_word_carries_what_would_help_answer_better() -> None:
    assert "Cần cột ngày" in _word_text(TRA_LOI)


# --- thu tu doc: canh bao TRUOC ket luan --------------------------------------


def test_the_warning_comes_before_the_first_claim_in_word() -> None:
    """Doc canh bao sau khi da doc het so la doc sau khi da tin."""
    text = _word_text(TRA_LOI)
    assert text.index("quá ít để nói gì") < text.index("Tư vấn tài chính")


def test_the_warning_sheet_is_the_first_one_in_excel() -> None:
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(to_excel(TRA_LOI)))
    assert book.worksheets[0].title == "Cảnh báo"


# --- cau tra loi rong --------------------------------------------------------


def test_an_answer_with_nothing_in_it_still_opens() -> None:
    # Mot tep hong con te hon mot tep rong: nguoi ta khong biet la khong co gi
    # hay la he thong loi.
    assert to_excel(TRONG)
    assert to_word(TRONG)


def test_an_empty_answer_says_so_rather_than_showing_a_blank_page() -> None:
    assert "Không có kết luận nào" in _word_text(TRONG)


def test_the_claims_sheet_exists_even_when_empty() -> None:
    assert SHEET_CLAIMS in _excel_text(TRONG)


# --- van la tep that ----------------------------------------------------------


def test_what_comes_out_is_a_real_xlsx() -> None:
    assert to_excel(TRA_LOI)[:2] == b"PK"


def test_what_comes_out_is_a_real_docx() -> None:
    assert to_word(TRA_LOI)[:2] == b"PK"


# --- cau tra loi thang phai di theo ra ngoai ----------------------------------

CO_CAU_CHOT = ManagerAnswer(
    question="Yếu tố nào ảnh hưởng mạnh hơn?",
    summary="poutcome ảnh hưởng mạnh hơn campaign tới tỷ lệ mở sổ.",
    claims=(
        ClaimEvidence(
            claim="Nhóm success đạt 65.11 %.",
            metric_keys=("y.yes.share_pct.by.poutcome.success",),
            evidence_ref="mart://x.parquet",
        ),
    ),
    warnings=("bỏ qua 1 nhóm có dưới 5 dòng, quá ít để nói gì",),
)


def test_excel_carries_the_direct_answer() -> None:
    """Tep nay di ra ngoai cho nguoi khong mo dashboard duoc, nen dong quan
    trong nhat khong duoc phep o lai tren man hinh."""
    assert "poutcome ảnh hưởng mạnh hơn" in _excel_text(CO_CAU_CHOT)


def test_word_carries_the_direct_answer() -> None:
    assert "poutcome ảnh hưởng mạnh hơn" in _word_text(CO_CAU_CHOT)


def test_the_direct_answer_comes_before_the_claims_in_word() -> None:
    text = _word_text(CO_CAU_CHOT)
    assert text.index("poutcome ảnh hưởng mạnh hơn") < text.index("Nhóm success")


def test_the_warning_still_comes_before_the_direct_answer_in_word() -> None:
    """Mot cau chot doc truoc khi biet du lieu mong la mot cau chot duoc tin nham."""
    text = _word_text(CO_CAU_CHOT)
    assert text.index("quá ít để nói gì") < text.index("poutcome ảnh hưởng mạnh hơn")


def test_an_answer_with_no_summary_still_opens() -> None:
    assert to_excel(TRA_LOI)[:2] == b"PK"
    assert to_word(TRA_LOI)[:2] == b"PK"
