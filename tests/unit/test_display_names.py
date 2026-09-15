"""Ten cot tieng Viet trong chu hien thi, va CHI trong chu hien thi.

Loi goc: cau tra loi thang va cac ket luan in nguyen ten cot tieng Anh cua du lieu
("ROA(A) before interest and % after tax") du bang chu giai da khai nghia tieng
Viet. Doi luc hien thi, bang code; metric key va phep tinh van dung ten goc.
"""

from __future__ import annotations

from io import BytesIO

from analysis_system.models.agents import ClaimEvidence, ManagerAnswer
from analysis_system.services.display_names import alias_of, column_aliases, localize
from analysis_system.services.export_answer import to_excel, to_word
from analysis_system.web.view import manager_answer

ALIASES = {
    "Net Income to Total Assets": "lợi nhuận ròng/tổng tài sản",
    "Total Assets": "tổng tài sản",
    "Debt Share %": "tỷ lệ nợ",
    "Age": "tuổi",
    "y": "đăng ký",
}


# --- doi ten trong mot cau ----------------------------------------------------------


def test_a_long_column_name_becomes_its_vietnamese_name() -> None:
    said = localize("Chỉ số là Net Income to Total Assets, cao nhất.", ALIASES)
    assert said == "Chỉ số là lợi nhuận ròng/tổng tài sản, cao nhất."


def test_a_name_that_starts_a_sentence_is_capitalised() -> None:
    said = localize(
        "Net Income to Total Assets có trung bình 0.81. Debt Share % thì thấp.", ALIASES
    )
    assert said == "Lợi nhuận ròng/tổng tài sản có trung bình 0.81. Tỷ lệ nợ thì thấp."


def test_the_longest_name_is_changed_first() -> None:
    """Khong ra "Net Income to tong tai san": ten ngan nam trong ten dai."""
    assert "Net Income" not in localize("So Net Income to Total Assets.", ALIASES)
    assert localize("Cả Total Assets.", ALIASES) == "Cả tổng tài sản."


def test_both_names_side_by_side_become_one() -> None:
    assert localize("Tỷ lệ nợ (Debt Share %) cao.", ALIASES) == "Tỷ lệ nợ cao."
    assert localize("Debt Share % (tỷ lệ nợ) cao.", ALIASES) == "Tỷ lệ nợ cao."


def test_a_name_inside_another_word_is_left_alone() -> None:
    assert (
        localize("Average và Agent không phải cột.", ALIASES) == "Average và Agent không phải cột."
    )


def test_a_short_name_changes_only_when_written_exactly() -> None:
    assert localize("Theo Age thì khác.", ALIASES) == "Theo tuổi thì khác."
    assert localize("Theo age thì khác.", ALIASES) == "Theo age thì khác."


def test_a_very_short_name_is_never_changed() -> None:
    assert localize("Nhóm y tế và y.", ALIASES) == "Nhóm y tế và y."


def test_a_long_name_matches_whatever_its_case() -> None:
    assert (
        localize("theo net income to total assets", ALIASES) == "theo lợi nhuận ròng/tổng tài sản"
    )


def test_no_glossary_changes_nothing() -> None:
    assert localize("Net Income to Total Assets", {}) == "Net Income to Total Assets"


def test_the_alias_is_the_first_way_of_saying_it() -> None:
    assert alias_of("tỷ lệ nợ; hệ số nợ") == "tỷ lệ nợ"
    assert column_aliases([(" Debt  Share %", "tỷ lệ nợ; hệ số nợ"), ("Khac", "")]) == {
        "Debt Share %": "tỷ lệ nợ"
    }


# --- trang va tep xuat -------------------------------------------------------------

TRA_LOI = ManagerAnswer(
    question="Ba chỉ số khác biệt nhất?",
    summary="Net Income to Total Assets khác biệt nhất.",
    claims=(
        ClaimEvidence(
            claim="Net Income to Total Assets có trung bình 0.81 ở nhóm không phá sản.",
            metric_keys=("Net Income to Total Assets.mean.by.flag.0",),
            evidence_ref="mart://bang.parquet",
        ),
    ),
    warnings=("Câu hỏi có nhắc tới cột Debt Share %, nhưng không kết luận nào dựa trên nó.",),
)


def test_the_page_shows_vietnamese_names_but_keeps_the_keys() -> None:
    payload = manager_answer(TRA_LOI, ALIASES)
    assert payload["summary"] == "Lợi nhuận ròng/tổng tài sản khác biệt nhất."
    assert payload["claims"][0]["claim"].startswith("Lợi nhuận ròng/tổng tài sản có trung bình")
    assert "tỷ lệ nợ" in payload["warnings"][0]
    # Khoa chi so giu ten goc: la thu de lan nguoc ve con so da tinh.
    assert payload["claims"][0]["metric_keys"] == ["Net Income to Total Assets.mean.by.flag.0"]


def test_without_a_glossary_the_page_is_unchanged() -> None:
    assert manager_answer(TRA_LOI)["summary"] == "Net Income to Total Assets khác biệt nhất."


def test_the_word_file_speaks_the_same_names() -> None:
    from docx import Document

    document = Document(BytesIO(to_word(TRA_LOI, ALIASES)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Lợi nhuận ròng/tổng tài sản khác biệt nhất." in text
    assert "Net Income to Total Assets.mean.by.flag.0" in text  # khoa van con


def test_the_excel_file_speaks_the_same_names() -> None:
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(to_excel(TRA_LOI, ALIASES)))
    cells = [
        str(cell)
        for sheet in book.worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell
    ]
    assert "Lợi nhuận ròng/tổng tài sản khác biệt nhất." in cells
    assert any("Net Income to Total Assets.mean.by.flag.0" in cell for cell in cells)
