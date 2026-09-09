"""Writing a table out in whatever format the reader asked for.

The format is the reader's choice and has nothing to do with what the data
arrived as. A PDF that was read into a table can leave as a spreadsheet; a
spreadsheet can leave as a Word document; a pile of text that was counted into
a table can leave as any of them. There is no fixed pairing of input to output
anywhere in here, and there should not be: the person asking is the one who
knows what they are going to do with it next.

Every writer takes the same two arguments and returns nothing, so adding a
format means adding one function and one line to the registry - and the CLI,
the help text and the error messages all pick it up without being told.

PDF is deliberately absent. The spec put it off rather than pull in a layout
engine for it, so asking for one gets that answer instead of "unknown format",
which would suggest a typo.
"""

from __future__ import annotations

import json
from collections.abc import Callable
from pathlib import Path
from typing import Final

import pandas as pd

from analysis_system.services import storage


class ExportError(RuntimeError):
    """A table cannot be written the way it was asked for."""


# Word puts every cell in its own XML element, so a big table stops being a
# document and becomes a hang. Refused rather than silently truncated: losing
# rows without saying so is worse than not writing the file.
DOCX_MAX_ROWS: Final[int] = 2000

# Named here rather than folded into the unknown-format message, so the answer
# says *why* and can be traced to the decision that caused it.
DEFERRED: Final[dict[str, str]] = {
    "pdf": (
        "Xuat PDF dang HOAN theo BUILD_SPEC muc 3 - chua them thu vien dan trang nao. "
        "Xuat ra 'html' roi in tu trinh duyet la duong ngan nhat hien co."
    )
}


def _to_csv(frame: pd.DataFrame, path: Path) -> None:
    frame.to_csv(path, index=False)


def _to_xlsx(frame: pd.DataFrame, path: Path) -> None:
    frame.to_excel(path, index=False, engine="openpyxl")


def _to_html(frame: pd.DataFrame, path: Path) -> None:
    path.write_text(frame.to_html(index=False, border=0), encoding="utf-8")


def _to_json(frame: pd.DataFrame, path: Path) -> None:
    # Records rather than pandas' default, because the reader of a .json export
    # is usually another program and a list of rows is what it expects.
    rows = json.loads(frame.to_json(orient="records", date_format="iso"))
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


def _cell(value: object) -> str:
    """One cell as text, with the pipes that would break a Markdown table."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ")


def _to_markdown(frame: pd.DataFrame, path: Path) -> None:
    # Written out by hand rather than through `to_markdown`, which needs
    # tabulate - a dependency outside the spec's list for one table border.
    columns = [_cell(name) for name in frame.columns]
    lines = ["| " + " | ".join(columns) + " |", "|" + "|".join(["---"] * len(columns)) + "|"]
    lines.extend(
        "| " + " | ".join(_cell(value) for value in row) + " |"
        for row in frame.itertuples(index=False, name=None)
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _to_docx(frame: pd.DataFrame, path: Path) -> None:
    if len(frame) > DOCX_MAX_ROWS:
        raise ExportError(
            f"Bang co {len(frame):,} dong, qua lon de dat vao mot file Word "
            f"(gioi han {DOCX_MAX_ROWS:,}). Dung --rows de lay phan dau, "
            "hoac xuat ra 'xlsx' - Excel chiu duoc so dong nay."
        )
    from docx import Document  # imported here so the other formats never need it

    document = Document()
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.paragraphs[0].add_run(_cell(name)).bold = True
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = _cell(value)
    # str(), because python-docx types its argument as a path string rather
    # than a Path even though it accepts one.
    document.save(str(path))


Writer = Callable[[pd.DataFrame, Path], None]

WRITERS: Final[dict[str, Writer]] = {
    "csv": _to_csv,
    "xlsx": _to_xlsx,
    "docx": _to_docx,
    "md": _to_markdown,
    "html": _to_html,
    "json": _to_json,
}

# What a file extension means, for the times somebody writes `--out bao_cao.docx`
# and expects that to be enough. It is.
SUFFIXES: Final[dict[str, str]] = {
    ".csv": "csv",
    ".xlsx": "xlsx",
    ".xls": "xlsx",
    ".docx": "docx",
    ".doc": "docx",
    ".md": "md",
    ".markdown": "md",
    ".html": "html",
    ".htm": "html",
    ".json": "json",
    ".pdf": "pdf",
}


def available() -> str:
    """The formats that can be written, for a message that has to list them."""
    return ", ".join(sorted(WRITERS))


def format_for(target: Path, requested: str = "") -> str:
    """Which format to write, from what was asked for or from the file name.

    An explicit request wins. Otherwise the extension decides, because writing
    `--out bao_cao.docx` and being handed a CSV would be its own small betrayal.

    Raises:
        ExportError: the format is unknown, deferred, or impossible to guess.
    """
    name = (requested or SUFFIXES.get(target.suffix.lower(), "")).strip().lower().lstrip(".")
    if not name:
        raise ExportError(
            f"Khong doan duoc dinh dang tu duoi file {target.suffix!r}. "
            f"Dat --dinh-dang, hoac dat ten file co duoi. Co: {available()}."
        )
    if name in DEFERRED:
        raise ExportError(DEFERRED[name])
    if name not in WRITERS:
        raise ExportError(f"Khong biet dinh dang {name!r}. Co: {available()}.")
    return name


def write(frame: pd.DataFrame, target: Path, requested: str = "") -> str:
    """Write a table out, and say which format was used.

    Raises:
        ExportError: the format is not one that can be written.
        OSError: the destination refused the write.
    """
    name = format_for(target, requested)
    target.parent.mkdir(parents=True, exist_ok=True)
    storage.ensure_writable_directory(target.parent)
    WRITERS[name](frame, target)
    return name
